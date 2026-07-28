import docx
from pathlib import Path
from typing import List, Dict, Any

from utils.helpers import clean_text

class DOCXParser:
    def __init__(self):
        pass

    def parse(self, file_path: Path) -> List[Dict[str, Any]]:
        """
        Extract text from a DOCX file.
        Preserves heading styles by formatting them as Markdown headers.
        """
        doc = docx.Document(str(file_path))
        extracted_paragraphs = []
        
        for p in doc.paragraphs:
            text = clean_text(p.text)
            if not text:
                continue
                
            style_name = p.style.name if p.style else ""
            
            # Preserve headings using Markdown syntax
            if style_name.startswith('Heading'):
                try:
                    level = int(style_name.split(' ')[-1])
                    prefix = '#' * level + ' '
                except ValueError:
                    prefix = '# '
                
                text = f"{prefix}{text}"
            
            extracted_paragraphs.append(text)

        # Also extract structured content from Word tables
        for table_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [clean_text(cell.text.strip()) for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            if table_rows:
                header_sep = " | ".join(["---"] * len(table_rows[0].split(" | ")))
                md_table = f"\n### Table {table_idx+1}\n" + table_rows[0] + "\n" + header_sep + "\n" + "\n".join(table_rows[1:]) + "\n"
                extracted_paragraphs.append(md_table)
            
        full_text = "\n\n".join(extracted_paragraphs)
        
        return [{
            "text": full_text,
            "source": "digital"
        }]

def parse_docx(file_path: Path) -> List[Dict[str, Any]]:
    parser = DOCXParser()
    return parser.parse(file_path)
